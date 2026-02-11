from pathlib import Path

from openpyxl import Workbook

from src.excel_import import parse_excel_preview


def test_parse_excel_preview_detects_sections_questions_and_trainees(tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Lecture de plan"
    ws2 = wb.create_sheet("Acteur de l'acte de construire")
    wb.create_sheet("Feuil5")

    ws["E4"] = "Dupont Alice"
    ws["H4"] = "Durand Bob"
    ws["B8"] = 1
    ws["C8"] = "A"
    ws["E8"] = "a"
    ws["H8"] = "b"

    ws2["E4"] = "='Lecture de plan'!E4"
    ws2["H4"] = "='Lecture de plan'!H4"
    ws2["B9"] = 2
    ws2["C9"] = "d"

    path = tmp_path / "import_test.xlsx"
    wb.save(path)

    preview = parse_excel_preview(str(path))

    assert preview.questionnaire_name == "import_test"
    assert preview.section_names == ["Lecture de plan", "Acteur de l'acte de construire"]
    assert preview.total_questions == 2
    assert "Dupont Alice" in preview.trainees
