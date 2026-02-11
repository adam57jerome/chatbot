from __future__ import annotations

from io import BytesIO
from typing import Dict, Iterable

import openpyxl
import pandas as pd

from .utils import LearnerColumn, ParsedWorkbook, SubjectSheet

SYNTHESIS_SHEETS = {"Feuil5", "Synthèse", "Synthese", "synthese", "synthèse"}


def _iter_subject_sheet_names(sheet_names: Iterable[str]) -> Iterable[str]:
    for sheet_name in sheet_names:
        if sheet_name not in SYNTHESIS_SHEETS:
            yield sheet_name


def _is_empty(value: object) -> bool:
    return value is None or str(value).strip() == ""


def detect_learners(ws) -> list[LearnerColumn]:
    learners: list[LearnerColumn] = []
    for col in range(4, ws.max_column + 1):
        header = ws.cell(row=4, column=col).value
        if _is_empty(header):
            continue

        name = str(header).strip()
        # Exclure les en-têtes techniques éventuels
        lowered = name.lower()
        if lowered.startswith("score") or lowered in {"question", "bonne réponse"}:
            continue

        learners.append(LearnerColumn(name=name, response_col=col, score_col=col + 1))
    return learners


def detect_question_rows(ws, learners: list[LearnerColumn], start_row: int = 8, max_empty_streak: int = 3) -> list[int]:
    rows: list[int] = []
    started = False
    empty_streak = 0

    for row in range(start_row, ws.max_row + 1):
        q_number = ws.cell(row=row, column=2).value
        correct_answer = ws.cell(row=row, column=3).value
        has_response = any(not _is_empty(ws.cell(row=row, column=l.response_col).value) for l in learners)

        row_has_data = (not _is_empty(q_number)) or (not _is_empty(correct_answer)) or has_response

        if row_has_data:
            rows.append(row)
            started = True
            empty_streak = 0
        elif started:
            empty_streak += 1
            if empty_streak >= max_empty_streak:
                break

    return rows


def parse_subject_sheet(ws, sheet_name: str) -> SubjectSheet:
    learners = detect_learners(ws)
    if not learners:
        raise ValueError("structure non reconnue: aucun apprenant détecté en ligne 4")

    question_rows = detect_question_rows(ws, learners=learners, start_row=8)
    if not question_rows:
        raise ValueError("structure non reconnue: aucune question détectée")

    rows = []
    for idx, row in enumerate(question_rows, start=1):
        q_value = ws.cell(row=row, column=2).value
        data: Dict[str, object] = {
            "question": idx if _is_empty(q_value) else q_value,
            "correct_answer": ws.cell(row=row, column=3).value,
            "excel_row": row,
        }
        for learner in learners:
            data[f"{learner.name}__response"] = ws.cell(row=row, column=learner.response_col).value
        rows.append(data)

    df = pd.DataFrame(rows)
    start_row = question_rows[0]
    end_row = question_rows[-1]
    total_row = end_row + 2
    note_row = end_row + 3
    return SubjectSheet(
        name=sheet_name,
        learners=learners,
        questions_df=df,
        question_start_row=start_row,
        question_end_row=end_row,
        total_row=total_row,
        note_row=note_row,
    )


def parse_workbook(file_bytes: bytes) -> ParsedWorkbook:
    wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=False)
    subjects: Dict[str, SubjectSheet] = {}
    ignored_sheets: Dict[str, str] = {}

    for sheet_name in _iter_subject_sheet_names(wb.sheetnames):
        ws = wb[sheet_name]
        try:
            subjects[sheet_name] = parse_subject_sheet(ws, sheet_name)
        except ValueError as exc:
            ignored_sheets[sheet_name] = str(exc)

    return ParsedWorkbook(subjects=subjects, ignored_sheets=ignored_sheets)
