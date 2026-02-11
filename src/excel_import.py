from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from .models import Answer, Attempt, Question, Questionnaire, Section, Session as CohortSession, Trainee
from .services import score_answer


@dataclass
class ExcelSectionData:
    name: str
    questions: list[tuple[int, str]]
    responses_by_trainee: dict[str, dict[int, str]]


@dataclass
class ExcelImportPreview:
    source_path: str
    questionnaire_name: str
    section_names: list[str]
    trainees: list[str]
    total_questions: int
    sections_data: list[ExcelSectionData]


def _normalize_spaces(value: str) -> str:
    return " ".join(value.strip().split())


def _formula_ref(value: object) -> tuple[str, str] | None:
    if not isinstance(value, str):
        return None
    text = value.strip()
    m = re.match(r"^=\s*'?(?P<sheet>[^']+)'?!(?P<cell>[A-Za-z]+\d+)$", text)
    if not m:
        return None
    return m.group("sheet"), m.group("cell").upper()


def _resolve_cell_value(wb_formula, wb_data, sheet_name: str, cell_ref: str) -> object:
    formula_sheet = wb_formula[sheet_name]
    v = formula_sheet[cell_ref].value
    fref = _formula_ref(v)
    if fref:
        ref_sheet, ref_cell = fref
        if ref_sheet in wb_data.sheetnames:
            fallback = wb_data[ref_sheet][ref_cell].value
            if fallback not in (None, ""):
                return fallback
        if ref_sheet in wb_formula.sheetnames:
            return wb_formula[ref_sheet][ref_cell].value
    if v not in (None, ""):
        return v
    if sheet_name in wb_data.sheetnames:
        return wb_data[sheet_name][cell_ref].value
    return None


def _to_int_or_none(v: object) -> int | None:
    if isinstance(v, int):
        return v
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str):
        s = v.strip()
        if s.isdigit():
            return int(s)
    return None


def _extract_trainees(wb_formula, wb_data, sheet_name: str) -> list[tuple[int, str]]:
    ws = wb_formula[sheet_name]
    trainees: list[tuple[int, str]] = []
    seen: set[str] = set()
    max_col = max(200, ws.max_column)
    for col in range(5, max_col + 1):
        cell_ref = f"{ws.cell(row=4, column=col).column_letter}4"
        raw = _resolve_cell_value(wb_formula, wb_data, sheet_name, cell_ref)
        if raw in (None, ""):
            continue
        name = _normalize_spaces(str(raw))
        if not name or name.lower() in {"nom", "stagiaire"}:
            continue
        if name in seen:
            continue
        seen.add(name)
        trainees.append((col, name))
    return trainees


def _extract_section(ws_formula, trainees_cols: list[tuple[int, str]]) -> ExcelSectionData:
    start_row: int | None = None
    for row in range(6, 201):
        maybe = _to_int_or_none(ws_formula.cell(row=row, column=2).value)
        if maybe is not None:
            start_row = row
            break

    questions: list[tuple[int, str]] = []
    responses_by_trainee: dict[str, dict[int, str]] = {name: {} for _, name in trainees_cols}
    if start_row is None:
        return ExcelSectionData(name=ws_formula.title, questions=questions, responses_by_trainee=responses_by_trainee)

    row = start_row
    while row <= 200:
        numero = _to_int_or_none(ws_formula.cell(row=row, column=2).value)
        if numero is None:
            break
        bonne = _normalize_spaces(str(ws_formula.cell(row=row, column=3).value or "")).lower()
        questions.append((numero, bonne))

        for col, trainee_name in trainees_cols:
            response = ws_formula.cell(row=row, column=col).value
            responses_by_trainee[trainee_name][numero] = _normalize_spaces(str(response or "")).lower()
        row += 1

    return ExcelSectionData(name=ws_formula.title, questions=questions, responses_by_trainee=responses_by_trainee)


def parse_excel_preview(file_path: str) -> ExcelImportPreview:
    wb_formula = load_workbook(file_path, data_only=False)
    wb_data = load_workbook(file_path, data_only=True)

    section_names = [name for name in wb_formula.sheetnames if name.lower() != "feuil5"]
    if not section_names:
        raise ValueError("Aucune feuille section détectée (toutes les feuilles étaient Feuil5 ?).")

    trainees_cols = _extract_trainees(wb_formula, wb_data, section_names[0])
    sections_data = []
    for section_name in section_names:
        ws = wb_formula[section_name]
        sections_data.append(_extract_section(ws, trainees_cols))

    total_questions = sum(len(section.questions) for section in sections_data)
    trainees = [name for _, name in trainees_cols]
    questionnaire_name = Path(file_path).stem
    return ExcelImportPreview(
        source_path=file_path,
        questionnaire_name=questionnaire_name,
        section_names=section_names,
        trainees=trainees,
        total_questions=total_questions,
        sections_data=sections_data,
    )


def _unique_questionnaire_name(db: Session, base_name: str) -> str:
    if not db.scalar(select(Questionnaire).where(Questionnaire.name == base_name).limit(1)):
        return base_name
    i = 2
    while True:
        candidate = f"{base_name} (import {i})"
        if not db.scalar(select(Questionnaire).where(Questionnaire.name == candidate).limit(1)):
            return candidate
        i += 1


def _split_name(fullname: str) -> tuple[str, str]:
    parts = fullname.split()
    if len(parts) <= 1:
        return fullname, ""
    return parts[0], " ".join(parts[1:])


def _get_or_create_session(db: Session, session_code: str) -> CohortSession:
    existing = db.scalar(select(CohortSession).where(CohortSession.code == session_code).limit(1))
    if existing:
        return existing
    session = CohortSession(code=session_code, label=f"Session {session_code}", active=True)
    db.add(session)
    db.commit()
    db.refresh(session)
    return session


def _get_or_create_trainee(db: Session, session_id: int, full_name: str) -> Trainee:
    nom, prenom = _split_name(full_name)
    existing = db.scalar(
        select(Trainee).where(
            Trainee.session_id == session_id,
            Trainee.nom == nom,
            Trainee.prenom == prenom,
        ).limit(1)
    )
    if existing:
        return existing
    trainee = Trainee(session_id=session_id, nom=nom, prenom=prenom, actif=True)
    db.add(trainee)
    db.commit()
    db.refresh(trainee)
    return trainee


def _get_or_create_attempt(db: Session, session_id: int, questionnaire_id: int, trainee_id: int) -> Attempt:
    existing = db.scalar(
        select(Attempt).where(
            Attempt.session_id == session_id,
            Attempt.questionnaire_id == questionnaire_id,
            Attempt.trainee_id == trainee_id,
        ).limit(1)
    )
    if existing:
        return existing
    attempt = Attempt(session_id=session_id, questionnaire_id=questionnaire_id, trainee_id=trainee_id)
    db.add(attempt)
    db.commit()
    db.refresh(attempt)
    return attempt


def import_excel_to_db(
    db: Session,
    preview: ExcelImportPreview,
    questionnaire_name: str,
    session_code: str,
    import_answers: bool,
) -> tuple[Questionnaire, CohortSession]:
    final_name = _unique_questionnaire_name(db, questionnaire_name)
    questionnaire = Questionnaire(name=final_name, description=f"Import Excel: {Path(preview.source_path).name}", active=True)
    db.add(questionnaire)
    db.flush()

    question_map: dict[int, Question] = {}
    for idx, section_data in enumerate(preview.sections_data, start=1):
        section = Section(questionnaire_id=questionnaire.id, name=section_data.name, ordre=idx)
        db.add(section)
        db.flush()
        for numero, bonne in section_data.questions:
            question = Question(
                questionnaire_id=questionnaire.id,
                section_id=section.id,
                numero=numero,
                bonne_reponse=bonne,
            )
            db.add(question)
            db.flush()
            question_map[numero] = question

    session = _get_or_create_session(db, session_code)

    if import_answers:
        for trainee_name in preview.trainees:
            trainee = _get_or_create_trainee(db, session.id, trainee_name)
            attempt = _get_or_create_attempt(db, session.id, questionnaire.id, trainee.id)
            for section_data in preview.sections_data:
                for numero, response in section_data.responses_by_trainee.get(trainee_name, {}).items():
                    question = question_map.get(numero)
                    if question is None:
                        continue
                    score = score_answer(response, question.bonne_reponse)
                    existing_answer = db.scalar(
                        select(Answer).where(Answer.attempt_id == attempt.id, Answer.question_id == question.id).limit(1)
                    )
                    if existing_answer:
                        existing_answer.reponse_texte = response
                        existing_answer.score = score
                    else:
                        db.add(
                            Answer(
                                attempt_id=attempt.id,
                                question_id=question.id,
                                reponse_texte=response,
                                score=score,
                            )
                        )

    db.commit()
    db.refresh(questionnaire)
    return questionnaire, session
