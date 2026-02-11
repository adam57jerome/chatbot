from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook


@dataclass
class ExcelQuestionRow:
    numero: int
    bonne_reponse_raw: str
    bonne_reponse_normalized: str | None
    row_index: int


@dataclass
class ExcelSectionData:
    name: str
    questions: list[ExcelQuestionRow]
    responses_by_trainee: dict[str, dict[int, str]]
    warnings: list[str] = field(default_factory=list)


@dataclass
class ExcelImportPreview:
    source_path: str
    questionnaire_name: str
    section_names: list[str]
    trainees: list[str]
    total_questions: int
    sections_data: list[ExcelSectionData]
    warnings: list[str] = field(default_factory=list)


def normalize_answer(value: object, ignore_accents: bool = True) -> str:
    import unicodedata

    text = " ".join(str(value or "").strip().lower().split())
    if ignore_accents:
        text = "".join(c for c in unicodedata.normalize("NFD", text) if unicodedata.category(c) != "Mn")
    return text


def normalize_good_answer(value: object) -> str | None:
    normalized = normalize_answer(value)
    return normalized if normalized in {"a", "b", "c", "d"} else None


def _formula_ref(value: object) -> tuple[str, str] | None:
    if not isinstance(value, str):
        return None
    text = value.strip()
    m = re.match(r"^=\s*'?(?P<sheet>[^']+)'?!(?P<cell>[A-Za-z]+\d+)$", text)
    if not m:
        return None
    return m.group("sheet"), m.group("cell").upper()


def _resolve_cell_value(wb_formula, wb_data, sheet_name: str, cell_ref: str) -> object:
    v = wb_formula[sheet_name][cell_ref].value
    fref = _formula_ref(v)
    if fref:
        ref_sheet, ref_cell = fref
        if ref_sheet in wb_data.sheetnames:
            data_value = wb_data[ref_sheet][ref_cell].value
            if data_value not in (None, ""):
                return data_value
        if ref_sheet in wb_formula.sheetnames:
            return wb_formula[ref_sheet][ref_cell].value
    if v not in (None, ""):
        return v
    if sheet_name in wb_data.sheetnames:
        return wb_data[sheet_name][cell_ref].value
    return None


def _to_int_or_none(v: object) -> int | None:
    if isinstance(v, int):
        return v
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str) and v.strip().isdigit():
        return int(v.strip())
    return None


def _extract_trainees(wb_formula, wb_data, sheet_name: str) -> list[tuple[int, str]]:
    ws = wb_formula[sheet_name]
    seen: set[str] = set()
    result: list[tuple[int, str]] = []

    for col in range(5, max(200, ws.max_column) + 1):
        cell_ref = f"{ws.cell(row=4, column=col).column_letter}4"
        raw = _resolve_cell_value(wb_formula, wb_data, sheet_name, cell_ref)
        name = normalize_answer(raw, ignore_accents=False)
        if not name or name in {"nom", "stagiaire"}:
            continue
        name = " ".join(str(raw).strip().split())
        if name in seen:
            continue
        seen.add(name)
        result.append((col, name))

    return result


def _extract_section(ws, trainees_cols: list[tuple[int, str]]) -> ExcelSectionData:
    start_row: int | None = None
    for row in range(6, 201):
        if _to_int_or_none(ws.cell(row=row, column=2).value) is not None:
            start_row = row
            break

    responses_by_trainee: dict[str, dict[int, str]] = {name: {} for _, name in trainees_cols}
    if start_row is None:
        return ExcelSectionData(name=ws.title, questions=[], responses_by_trainee=responses_by_trainee)

    questions: list[ExcelQuestionRow] = []
    warnings: list[str] = []

    row = start_row
    while row <= 200:
        numero = _to_int_or_none(ws.cell(row=row, column=2).value)
        if numero is None:
            break

        raw_good = " ".join(str(ws.cell(row=row, column=3).value or "").strip().split())
        normalized_good = normalize_good_answer(raw_good)
        if normalized_good is None:
            warnings.append(f"Section '{ws.title}' Q{numero}: clé de réponse manquante/invalide ('{raw_good}').")

        questions.append(
            ExcelQuestionRow(
                numero=numero,
                bonne_reponse_raw=raw_good,
                bonne_reponse_normalized=normalized_good,
                row_index=row,
            )
        )

        for col, trainee_name in trainees_cols:
            response_raw = ws.cell(row=row, column=col).value
            responses_by_trainee[trainee_name][numero] = " ".join(str(response_raw or "").strip().split())

        row += 1

    return ExcelSectionData(
        name=ws.title,
        questions=questions,
        responses_by_trainee=responses_by_trainee,
        warnings=warnings,
    )


def parse_excel_preview(file_path: str) -> ExcelImportPreview:
    wb_formula = load_workbook(file_path, data_only=False)
    wb_data = load_workbook(file_path, data_only=True)

    section_names = [name for name in wb_formula.sheetnames if name.lower() != "feuil5"]
    if not section_names:
        raise ValueError("Aucune feuille section détectée.")

    trainees_cols = _extract_trainees(wb_formula, wb_data, section_names[0])
    sections_data = [_extract_section(wb_formula[name], trainees_cols) for name in section_names]

    warnings = [w for section in sections_data for w in section.warnings]

    return ExcelImportPreview(
        source_path=file_path,
        questionnaire_name=Path(file_path).stem,
        section_names=section_names,
        trainees=[name for _, name in trainees_cols],
        total_questions=sum(len(section.questions) for section in sections_data),
        sections_data=sections_data,
        warnings=warnings,
    )
