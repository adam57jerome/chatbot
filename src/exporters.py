from __future__ import annotations

from io import BytesIO

import openpyxl
import pandas as pd

from .utils import SubjectSheet


def _write_subject_sheet(ws, subject: SubjectSheet) -> None:
    ws.cell(row=4, column=2, value="Question")
    ws.cell(row=4, column=3, value="Bonne réponse")
    for learner in subject.learners:
        ws.cell(row=4, column=learner.response_col, value=learner.name)
        ws.cell(row=4, column=learner.score_col, value=f"Score {learner.name}")

    for _, row in subject.questions_df.iterrows():
        excel_row = int(row["excel_row"])
        ws.cell(row=excel_row, column=2, value=row["question"])
        ws.cell(row=excel_row, column=3, value=row["correct_answer"])
        for learner in subject.learners:
            response_col = f"{learner.name}__response"
            score_col = f"{learner.name}__score"
            ws.cell(row=excel_row, column=learner.response_col, value=row.get(response_col))
            ws.cell(row=excel_row, column=learner.score_col, value=int(row.get(score_col, 0)))

        for learner in subject.learners:
            ws.cell(row=subject.total_row, column=learner.score_col, value=subject.totals.get(learner.name, 0))
            ws.cell(row=subject.note_row, column=learner.score_col, value=subject.notes.get(learner.name, 0.0))

    ws.cell(row=subject.total_row, column=3, value="Total points")
    ws.cell(row=subject.note_row, column=3, value="Note /20")


def _write_synthesis_sheet(wb, synthesis_df: pd.DataFrame) -> None:
    if "Feuil5" in wb.sheetnames:
        del wb["Feuil5"]
    if "Synthèse" in wb.sheetnames:
        del wb["Synthèse"]

    ws_syn = wb.create_sheet("Feuil5")
    for col_idx, col_name in enumerate(synthesis_df.columns, start=1):
        ws_syn.cell(row=1, column=col_idx, value=col_name)
    for row_idx, row in enumerate(synthesis_df.itertuples(index=False), start=2):
        for col_idx, value in enumerate(row, start=1):
            ws_syn.cell(row=row_idx, column=col_idx, value=value)


def export_recalculated_workbook(
    source_bytes: bytes,
    subjects: dict[str, SubjectSheet],
    synthesis_df: pd.DataFrame,
) -> bytes:
    wb = openpyxl.load_workbook(BytesIO(source_bytes), data_only=False)

    for subject_name, subject in subjects.items():
        if subject_name not in wb.sheetnames:
            continue
        ws = wb[subject_name]
        _write_subject_sheet(ws, subject)

    _write_synthesis_sheet(wb, synthesis_df)

    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return out.getvalue()


def export_manual_workbook(
    section_name: str,
    subjects: dict[str, SubjectSheet],
    synthesis_df: pd.DataFrame,
) -> bytes:
    wb = openpyxl.Workbook()
    default = wb.active
    wb.remove(default)

    for subject_name, subject in subjects.items():
        ws = wb.create_sheet(subject_name[:31])
        ws.cell(row=1, column=1, value="Section")
        ws.cell(row=1, column=2, value=section_name)
        _write_subject_sheet(ws, subject)

    _write_synthesis_sheet(wb, synthesis_df)

    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return out.getvalue()


def export_synthesis_csv(synthesis_df: pd.DataFrame) -> bytes:
    return synthesis_df.to_csv(index=False).encode("utf-8")
